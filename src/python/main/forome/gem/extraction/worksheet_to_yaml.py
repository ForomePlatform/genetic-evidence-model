#!/usr/bin/env python3
"""worksheet_to_yaml.py - deterministic conversion of a GEM annotation worksheet
(GEM_annotation_worksheet.xlsx) into a corpus annotation YAML.

The conversion is mechanical: every value in the output was typed or selected by
the annotator in the worksheet; nothing is inferred, normalised, reworded or
filled in (one provenance exception: a blank annotation_finished defaults to
the conversion date, recorded as annotated_at_defaulted). Enumerated values
are checked against the workbook's own Enums sheet
(the single source of truth shared with the dropdowns). On any error the report
is printed and no YAML is written (unless --force).

Reads worksheet version 2 (Paper tab: pass_code, protocol_version, worksheet_version;
Effort tab: the EVALUATION_PROTOCOL section 10 fields).

Usage:
    gem-worksheet-to-yaml worksheet.xlsx annotation.yaml [--effort effort.csv] [--force]
    python3 -m forome.gem.extraction.worksheet_to_yaml ...   (same arguments)

Dependencies: openpyxl, via the [xl] extra (pip install 'gem-mapping-studio[xl]',
or pip install -e '.[xl]' in a checkout); pyyaml is core.
"""
import argparse
import csv
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

import yaml

MULTI = {
    "knowledge_domain": ("knowledge_domain_1", "knowledge_domain_2"),
    "method": ("method_1", "method_2"),
    "variant_ascertainment": ("variant_ascertainment_1", "variant_ascertainment_2"),
    "measurement_target": ("measurement_target_1", "measurement_target_2"),
}
ENUM_OF = {
    "knowledge_domain_1": "KNOWLEDGE_DOMAIN", "knowledge_domain_2": "KNOWLEDGE_DOMAIN",
    "method_1": "METHOD", "method_2": "METHOD", "primary_method": "METHOD",
    "target_type": "TARGET_TYPE", "resolution": "RESOLUTION",
    "phenotype_scale": "PHENOTYPE_SCALE", "credibility": "CREDIBILITY",
    "knockout_type": "KNOCKOUT_TYPE", "subdomain": "SUBDOMAIN",
    "penetrance": "PENETRANCE",
    "variant_ascertainment_1": "VARIANT_ASCERTAINMENT",
    "variant_ascertainment_2": "VARIANT_ASCERTAINMENT",
    "measurement_target_1": "MEASUREMENT_TARGET", "measurement_target_2": "MEASUREMENT_TARGET",
    "gene_relation": "GENE_RELATION",
}
BOOL_COLS = {"mendelian_segregation", "exact_variant", "genetic_background_considered",
             "cited_evidence"}
FREE_SINGLE = ["primary_method", "method_detail", "credibility_comment", "organism",
               "specificity_of_phenotype", "knockout_type", "mode_of_inheritance",
               "subdomain", "penetrance", "environmental_factors", "gene_relation",
               "gene_relation_detail"]
ALWAYS_REQUIRED = ["label", "knowledge_domain_1", "method_1", "target_type", "target",
                   "resolution", "phenotype_scale", "credibility"]
# Identifier columns are text. Google Sheets exports a typed 1 as the number 1.0,
# so without this an id typed as 1 on one sheet never matches a reference to it.
ID_COLS = {"item_id", "ce_id"}
CONVERTER_VERSION = "2.0"
# Worksheet revisions this converter reads (the Paper tab's worksheet_version).
SUPPORTED_WORKSHEET_VERSIONS = {"2"}
# EVALUATION_PROTOCOL section 8: manual pass codes and the output each produces.
PASS_OUTPUT = {"H0": "h0", "SR": "h1", "WR": "h0r"}
# EVALUATION_PROTOCOL section 10: one Effort row per sitting.
EFFORT_COLS = ["date", "activity", "start_time", "stop_time", "minutes", "interruptions",
               "external_resources", "technical_problems", "comment"]


def cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, bool):
        return v
    # Sheets exports date cells as midnight datetimes; the sheet asks for dates.
    if isinstance(v, datetime) and v.time() == datetime.min.time():
        return v.date()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def fmt(v):
    """CSV form of a cell: times as HH:MM (with seconds only when present)."""
    if isinstance(v, time):
        return v.strftime("%H:%M" if v.second == 0 else "%H:%M:%S")
    return v


def rows_as_dicts(ws):
    headers = [cell(c.value) for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        d = {h: cell(v) for h, v in zip(headers, r) if h}
        for h in ID_COLS & d.keys():
            if d[h] is not None:
                d[h] = str(d[h])
        if any(v is not None for k, v in d.items() if k != "word_count"):
            out.append(d)
    return out


def to_bool(v, where, errors):
    if v is None:
        return None
    s = str(v).strip().upper()
    if s in ("TRUE", "YES", "1"):
        return True
    if s in ("FALSE", "NO", "0"):
        return False
    errors.append(f"{where}: '{v}' is not TRUE/FALSE")
    return None


def semicolon_list(v):
    return [p.strip() for p in str(v).split(";") if p.strip()] if v else None


def id_list(v):
    """A list of ids: ';' or ',' separated (ids contain neither)."""
    return [p.strip() for p in re.split(r"[;,]", str(v)) if p.strip()] if v else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("worksheet")
    ap.add_argument("output")
    ap.add_argument("--effort", help="also write the Effort sheet to this CSV")
    ap.add_argument("--force", action="store_true", help="write YAML despite errors")
    args = ap.parse_args()

    # openpyxl is an optional extra: import on demand with an actionable error.
    try:
        import openpyxl
    except ImportError:
        sys.exit("The worksheet converter needs openpyxl. Install it with:  "
                 "pip install 'gem-mapping-studio[xl]'   or, in a checkout,  "
                 "pip install -e '.[xl]'")

    wb = openpyxl.load_workbook(args.worksheet, data_only=True)
    errors, warnings = [], []

    # Enums from the workbook itself: dropdowns and converter cannot drift apart.
    enums = {}
    ews = wb["Enums"]
    for j in range(1, ews.max_column + 1):
        name = cell(ews.cell(row=1, column=j).value)
        if name:
            enums[name] = [cell(ews.cell(row=i, column=j).value)
                           for i in range(2, ews.max_row + 1)
                           if cell(ews.cell(row=i, column=j).value)]

    # Paper sheet: field name in column A, value in column B.
    paper = {}
    for r in wb["Paper"].iter_rows(values_only=True):
        k, v = cell(r[0]), cell(r[1]) if len(r) > 1 else None
        if k:
            paper[k] = v
    for f in ("paper_id", "citation", "title", "annotator"):
        if not paper.get(f):
            errors.append(f"Paper sheet: '{f}' is empty")
    # Worksheet version 2 fields.
    wv = paper.get("worksheet_version")
    if wv is None:
        errors.append("Paper sheet: 'worksheet_version' is empty: this worksheet predates version 2 "
                      "and must be migrated (docs/WORKSHEET_SOP.md, 'Worksheet versions')")
    elif str(wv) not in SUPPORTED_WORKSHEET_VERSIONS:
        errors.append(f"Paper sheet: worksheet_version '{wv}' is not supported (this converter "
                      f"reads version {', '.join(sorted(SUPPORTED_WORKSHEET_VERSIONS))})")
    # protocol_version is sealed into the provenance, so anything not plainly a version
    # (an int such as 1 for "1.0", a date from locale parsing) is an error, not a guess.
    pv = paper.get("protocol_version")
    if pv is None:
        errors.append("Paper sheet: 'protocol_version' is empty")
    else:
        if isinstance(pv, float):  # unambiguous, e.g. 1.1; but the cell should be text
            warnings.append(f"Paper sheet: protocol_version is stored as a number, recorded as "
                            f"'{pv}'; format the cell as plain text")
            pv = str(pv)
        if not (isinstance(pv, str) and re.fullmatch(r"\d+(\.\d+)+", pv)):
            errors.append(f"Paper sheet: protocol_version {pv!r} is not a version such as 1.1; "
                          "format the cell as plain text and retype it")
    pc = paper.get("pass_code")
    if pc is None:
        errors.append("Paper sheet: 'pass_code' is empty")
    elif pc not in PASS_OUTPUT:  # study metadata fixed by the protocol, not a workbook enum
        errors.append(f"Paper sheet: pass_code '{pc}' is not one of {', '.join(PASS_OUTPUT)}")
    # The study layout is study/processed/<paper-id>/<pass>/<pass>.yaml: a sheet copied
    # from an earlier pass that still carries the old pass_code or paper_id is caught here.
    out = Path(args.output)
    if pc in PASS_OUTPUT and out.stem in PASS_OUTPUT.values() and out.stem != PASS_OUTPUT[pc]:
        errors.append(f"Paper sheet: pass_code {pc} produces {PASS_OUTPUT[pc]}.yaml, but the output "
                      f"is {out.name}; was the sheet copied from another pass?")
    if (out.parent.name == out.stem and paper.get("paper_id")
            and out.parent.parent.name != str(paper["paper_id"])):
        errors.append(f"output directory '{out.parent.parent.name}' does not match "
                      f"paper_id '{paper['paper_id']}'")
    missing = [h for h in EFFORT_COLS if h not in {cell(c.value) for c in wb["Effort"][1]}]
    if missing:
        errors.append(f"Effort sheet: missing column(s) {', '.join(missing)}")

    # Effort sittings outside annotation_started..today: a sheet copied from an earlier
    # pass still carrying its sittings, or a mistyped date. Advisory: rows are kept as typed.
    def day(v):  # datetime subclasses date, and the two cannot be compared
        return v.date() if isinstance(v, datetime) else v if isinstance(v, date) else None
    started = day(paper.get("annotation_started"))
    for d in rows_as_dicts(wb["Effort"]):
        when = day(d.get("date"))
        if when is not None:
            if started is not None and when < started:
                warnings.append(f"Effort sheet: sitting dated {when} is before annotation_started "
                                f"{started}; copied from an earlier pass?")
            elif when > date.today():
                warnings.append(f"Effort sheet: sitting dated {when} is in the future")

    annotated_at_defaulted = not paper.get("annotation_finished")
    if annotated_at_defaulted:
        paper["annotation_finished"] = date.today()
        warnings.append("Paper sheet: 'annotation_finished' is empty; "
                        "annotated_at defaults to the conversion date")

    items = rows_as_dicts(wb["Items"])
    if not items:
        errors.append("Items sheet: no items")
    asserts = rows_as_dicts(wb["Assertions"])
    ces = rows_as_dicts(wb["Candidate Extensions"])
    omitted = rows_as_dicts(wb["Omitted Dimensions"])

    # Enumeration and requiredness checks.
    for d in items:
        iid = d.get("item_id") or "item with empty item_id"
        for col, en in ENUM_OF.items():
            v = d.get(col)
            if v is not None and v not in enums.get(en, []):
                errors.append(f"{iid}.{col}: '{v}' not in {en}")
        for col in ALWAYS_REQUIRED:
            if d.get(col) is None:
                errors.append(f"{iid}: required '{col}' is empty")
        for col in BOOL_COLS:
            d[col] = to_bool(d.get(col), f"{iid}.{col}", errors)
        # Conditional-activation warnings (advisory, mirrors the SHACL rules).
        kd = [d.get("knowledge_domain_1"), d.get("knowledge_domain_2")]
        if d.get("knockout_type") and "MODEL_ORGANISM" not in kd:
            warnings.append(f"{iid}: knockout_type set but knowledge_domain has no MODEL_ORGANISM")
        if d.get("target_type") == "VARIANT" and not d.get("variant_ascertainment_1"):
            warnings.append(f"{iid}: target_type VARIANT but variant_ascertainment is empty")
        if (d.get("measurement_target_1") or d.get("gene_relation")) and "GENE_FUNCTION" not in kd:
            warnings.append(f"{iid}: measurement_target/gene_relation set without GENE_FUNCTION domain")
        if "IN_VIVO" in (d.get("method_1"), d.get("method_2")) or "MODEL_ORGANISM" in kd:
            if not d.get("organism"):
                warnings.append(f"{iid}: organism expected (IN_VIVO method or MODEL_ORGANISM domain)")

    item_ids = [d.get("item_id") for d in items]
    if len(item_ids) != len(set(item_ids)):
        errors.append("Items sheet: duplicate item_id")

    by_item = {}
    for d in asserts:
        iid, num = d.get("item_id"), d.get("assertion_number")
        where = f"assertion ({iid}, #{num})"
        if iid not in item_ids:
            errors.append(f"{where}: item_id not on Items sheet")
            continue
        for f in ("assertion_number", "statement", "page", "key_phrase"):
            if d.get(f) is None:
                errors.append(f"{where}: '{f}' is empty")
        # Same Sheets typing issue as the ids: a number typed as text arrives as str.
        if isinstance(num, str):
            if num.isdigit():
                d["assertion_number"] = int(num)
            else:
                errors.append(f"{where}: assertion_number '{num}' is not a whole number")
                d["assertion_number"] = None
        kp = d.get("key_phrase")
        if kp and len(str(kp).split()) > 15:
            warnings.append(f"{where}: key_phrase is {len(str(kp).split())} words (limit 15)")
        by_item.setdefault(iid, []).append(d)
    for iid in item_ids:
        if iid not in by_item:
            errors.append(f"{iid}: no assertions")

    ce_ids = [d.get("ce_id") for d in ces]
    present = [x for x in ce_ids if x]
    if len(present) != len(set(present)):
        errors.append("Candidate Extensions sheet: duplicate ce_id")
    for x in ce_ids:
        if x and not re.fullmatch(r"CE-\w+", x):
            warnings.append(f"Candidate Extensions: ce_id '{x}' is not in the CE-<n> form "
                            "the sheet asks for (kept as typed)")
    for d in items:
        for ref in id_list(d.get("candidate_extension_refs")) or []:
            if ref not in ce_ids:
                errors.append(f"{d.get('item_id')}: candidate_extension_refs '{ref}' not on Candidate Extensions sheet")

    # ---- Assemble the YAML structure (corpus shape) ----
    doc = {}
    doc["publication"] = {
        "id": f"PMID:{paper.get('pmid')}" if paper.get("pmid") else None,
        "doi": paper.get("doi"),
        "citation": paper.get("citation"),
        "title": paper.get("title"),
    }
    doc["publication"] = {k: v for k, v in doc["publication"].items() if v is not None}
    doc["provenance"] = {
        "mode": "manual",
        "pass_code": pc,
        "protocol_version": pv,
        "annotator": paper.get("annotator"),
        "annotated_at": str(paper.get("annotation_finished") or ""),
        "annotation_started": str(paper.get("annotation_started") or ""),
        "source": "manual_worksheet",
        "conversion": {
            "tool": "worksheet_to_yaml.py",
            "version": CONVERTER_VERSION,
            "deterministic": True,
            "converted_at": date.today().isoformat(),
            "worksheet_version": str(wv) if wv is not None else None,
            **({"annotated_at_defaulted": True} if annotated_at_defaulted else {}),
            "note": "Mechanical serialisation of the annotator's worksheet; no content decisions made in conversion.",
        },
    }
    if ces:
        doc["candidate_extensions"] = [{
            "id": d.get("ce_id"),
            "dimension": d.get("dimension_or_area"),
            "description": d.get("description"),
            "location_in_this_annotation": d.get("items_where_encountered"),
        } for d in ces]

    evidence = []
    for d in items:
        iid = d.get("item_id")
        e = {"id": iid, "label": d.get("label"), "source": "manual_annotation"}
        for out_name, cols in MULTI.items():
            vals = [d.get(c) for c in cols if d.get(c)]
            if vals:
                e[out_name] = vals
        e["target_type"] = d.get("target_type")
        e["target"] = d.get("target")
        e["resolution"] = d.get("resolution")
        e["phenotype_scale"] = d.get("phenotype_scale")
        e["credibility"] = d.get("credibility")
        for col in FREE_SINGLE:
            if d.get(col) is not None:
                e[col] = d.get(col)
        for col in BOOL_COLS:
            if d.get(col) is not None:
                e[col] = d.get(col)
        if d.get("hpo_terms"):
            e["hpo_terms"] = semicolon_list(d.get("hpo_terms"))
        sc = []
        for pair in semicolon_list(d.get("special_considerations")) or []:
            if ":" in pair:
                k, v = pair.split(":", 1)
                sc.append({"key": k.strip(), "value": v.strip()})
            else:
                sc.append({"key": "note", "value": pair})
        if sc:
            e["special_considerations"] = sc
        if d.get("notes_uncertainties"):
            sc.append({"key": "annotator_uncertainty", "value": d.get("notes_uncertainties")})
            e["special_considerations"] = sc
        refs = id_list(d.get("candidate_extension_refs"))
        if refs:
            e["candidate_extension_exercised"] = refs[0]
            if len(refs) > 1:
                e["candidate_extension_note"] = "Also exercises: " + ", ".join(refs[1:])
        e["assertions"] = []
        for a in sorted(by_item.get(iid, []), key=lambda x: (x.get("assertion_number") or 0)):
            aa = {
                "id": f"{iid}.A{a.get('assertion_number')}",
                "statement": a.get("statement"),
                "source_span": {"page": a.get("page"), "key_phrase": a.get("key_phrase")},
                "callout": ({"page": a.get("page"), "text": a.get("callout_note")}
                            if a.get("callout_note") else None),
            }
            e["assertions"].append(aa)
        evidence.append(e)
    doc["evidence"] = evidence

    if omitted:
        doc["annotator_omitted_dimensions"] = [{
            "dimension": d.get("dimension"), "rationale": d.get("rationale"),
        } for d in omitted]

    # ---- Report and write ----
    for w in warnings:
        print(f"WARNING: {w}")
    for e in errors:
        print(f"ERROR: {e}")
    print(f"{len(items)} items, {len(asserts)} assertions, {len(ces)} candidate extensions; "
          f"{len(errors)} errors, {len(warnings)} warnings.")
    if errors and not args.force:
        print("No YAML written. Fix the worksheet (or re-run with --force).")
        return 1

    class Dumper(yaml.SafeDumper):
        pass

    def str_presenter(dumper, data):
        if "\n" in data or len(data) > 90:
            return dumper.represent_scalar("tag:yaml.org,2002:str", data, style=">")
        return dumper.represent_scalar("tag:yaml.org,2002:str", data)

    Dumper.add_representer(str, str_presenter)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    with open(args.output, "w", encoding="utf-8", newline="\n") as f:
        f.write(f"# Annotation of {paper.get('paper_id')} - generated by worksheet_to_yaml.py "
                f"v{CONVERTER_VERSION} from the annotator's worksheet.\n"
                f"# Every content decision is the annotator's; the conversion is mechanical.\n")
        yaml.dump(doc, f, Dumper=Dumper, sort_keys=False, allow_unicode=True, width=80)
    print(f"Wrote {args.output}")

    if args.effort:
        eff = rows_as_dicts(wb["Effort"])
        Path(args.effort).parent.mkdir(parents=True, exist_ok=True)
        with open(args.effort, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["paper_id", "pass_code", "annotator"] + EFFORT_COLS)
            for d in eff:
                w.writerow([paper.get("paper_id"), pc, paper.get("annotator")]
                           + [fmt(d.get(c)) for c in EFFORT_COLS])
        print(f"Wrote {args.effort} ({len(eff)} sittings)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
